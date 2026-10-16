"""Utilidades de lectura y limpieza compartidas por los parsers."""
from __future__ import annotations

import csv
import hashlib
import re
from collections import Counter
from datetime import date, datetime
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from core.categorizer import normalize
from core.i18n import tr

CSV_ENCODINGS = ("utf-8-sig", "cp1252")
_MINUS = str.maketrans({"−": "-", "–": "-", " ": " "})


class ParseError(ValueError):
    pass


# ---------------------------------------------------------------- lectura de ficheros
_ZIP_MAGIC = b"PK\x03\x04"
_OLE_MAGIC = b"\xd0\xcf\x11\xe0"          # el .xls clasico (BIFF)


def load_rows(path: Path) -> list[list[Any]]:
    """Lee CSV/TSV/TXT o Excel como lista de filas crudas (sin asumir cabecera).

    El formato se decide por el contenido y no por la extension: muchos bancos (ING entre ellos) llaman `.xls` a un
    `.xlsx`, a un Excel antiguo de verdad o a una tabla HTML."""
    raw = path.read_bytes()
    if raw[:4] == _ZIP_MAGIC:
        return _xlsx_rows(raw, path.name)
    if raw[:4] == _OLE_MAGIC:
        return _xls_rows(path)
    if path.suffix.lower() in {".xlsx", ".xlsm", ".xls"} and raw.lstrip()[:1] == b"<":
        return _html_rows(raw, path.name)
    for enc in CSV_ENCODINGS:
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - cp1252 casi nunca falla
        raise ParseError(tr("No se pudo decodificar {name}").format(name=path.name))
    lines = text.splitlines()
    sample = "\n".join(lines[:20])
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=";,\t").delimiter
    except csv.Error:
        delimiter = ";"
    return [row for row in csv.reader(lines, delimiter=delimiter)]


def _xlsx_rows(raw: bytes, name: str) -> list[list[Any]]:
    from io import BytesIO

    from openpyxl import load_workbook          # solo se carga si se importa un Excel

    try:
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)       # desde memoria: openpyxl rechaza la extension .xls
    except Exception as exc:                    # archivo danado, protegido con contrasena...
        raise ParseError(tr("No se pudo abrir el Excel {name}: {exc}").format(name=name, exc=exc)) from exc
    try:
        for sheet in workbook.worksheets:       # la primera hoja con datos (algunos bancos dejan una portada vacia)
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            if any(any(cell not in (None, "") for cell in row) for row in rows):
                return rows
        return []
    finally:
        workbook.close()


def _xls_rows(path: Path) -> list[list[Any]]:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover - xlrd va en requirements.txt
        raise ParseError(tr("Falta el componente para leer archivos .xls antiguos.")) from exc
    try:
        book = xlrd.open_workbook(str(path))
    except Exception as exc:
        raise ParseError(tr("No se pudo abrir el Excel {name}: {exc}").format(name=path.name, exc=exc)) from exc
    for sheet in book.sheets():
        rows: list[list[Any]] = []
        for r in range(sheet.nrows):
            row: list[Any] = []
            for cell in sheet.row(r):
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                else:
                    row.append(cell.value)
            rows.append(row)
        if any(any(cell not in (None, "") for cell in row) for row in rows):
            return rows
    return []


class _TableReader(HTMLParser):
    """Las filas de la tabla mas grande de una pagina HTML (los 'Excel' que en realidad son HTML)."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._depth = 0
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag == "table":
            self._depth += 1
            self.tables.append([])
        elif tag == "tr" and self.tables:
            self.tables[-1].append([])
        elif tag in ("td", "th") and self.tables and self.tables[-1]:
            self._cell = []

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self._cell is not None:
            self.tables[-1][-1].append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "table":
            self._depth -= 1

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)


def _html_rows(raw: bytes, name: str) -> list[list[Any]]:
    for enc in CSV_ENCODINGS:
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover
        raise ParseError(tr("No se pudo decodificar {name}").format(name=name))
    reader = _TableReader()
    reader.feed(text)
    tables = [t for t in reader.tables if t]
    if not tables:
        raise ParseError(tr("No se encontró ninguna tabla en {name}.").format(name=name))
    return [list(row) for row in max(tables, key=len)]


def normalize_header(h: Any) -> str:
    """'Importe (EUR)' -> 'importe'; 'Fecha Operación' -> 'fecha operacion'."""
    text = re.sub(r"\(.*?\)", "", str(h or ""))
    return normalize(text).lower()


@dataclass(slots=True)
class Table:
    """Tabla minima: nombres de columna normalizados y una fila (dict) por registro."""

    columns: list[str]
    rows: list[dict[str, Any]]


def build_table(rows: list[list[Any]], header_index: int) -> Table:
    header = [normalize_header(h) for h in rows[header_index]]
    width = len(header)
    body = [r for r in rows[header_index + 1:] if any(str(c).strip() for c in r if c is not None)]
    return Table(header, [dict(zip(header, (list(r) + [None] * width)[:width])) for r in body])


# ---------------------------------------------------------------- conversiones
def to_decimal(value: Any) -> Decimal:
    """Acepta numeros de Excel o cadenas espanolas ('-2,99EUR', '1.234,56 €').

    En cadenas, '.' es separador de miles y ',' el decimal.
    """
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        raise ParseError(tr("Importe no valido: {value}").format(value=repr(value)))
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).translate(_MINUS)
    text = re.sub(r"[^\d,.\-+]", "", text)
    text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ParseError(tr("Importe no valido: {value}").format(value=repr(value))) from exc


def to_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if len(text) >= 10:                              # camino rapido para los formatos habituales (strptime es muy lento)
        try:
            if text[2] in "/-" and text[5] == text[2]:
                return date(int(text[6:10]), int(text[3:5]), int(text[:2]))
            if text[4] == "-" and text[7] == "-":
                return date(int(text[:4]), int(text[5:7]), int(text[8:10]))
        except ValueError:
            pass
    for fmt, size in (("%d/%m/%Y", 10), ("%Y-%m-%d", 10), ("%d-%m-%Y", 10)):
        try:
            return datetime.strptime(text[:size], fmt).date()
        except ValueError:
            continue
    raise ParseError(tr("Fecha no valida: {value}").format(value=repr(value)))


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


# ---------------------------------------------------------------- hashes anti-duplicado
def make_hash(*parts: Any) -> str:
    payload = "\x1f".join(str(p) for p in parts)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


class HashFactory:
    """Genera hashes estables. Si dos filas son identicas (misma fecha, concepto,
    importe y saldo) las distingue por orden de aparicion, de modo que dos pagos
    legitimos iguales no se pierden, y reimportar el mismo extracto da los mismos
    hashes."""

    def __init__(self, namespace: str) -> None:
        self._ns = namespace
        self._seen: Counter[str] = Counter()

    def make(self, *parts: Any) -> str:
        base = make_hash(self._ns, *parts)
        occurrence = self._seen[base]
        self._seen[base] += 1
        return make_hash(base, occurrence)
