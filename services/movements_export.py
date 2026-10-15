"""Exportar los movimientos que se ven en pantalla a CSV o Excel (nunca PDF)."""
from __future__ import annotations

import csv
from collections.abc import Mapping, Sequence
from decimal import Decimal
from pathlib import Path


from core.entities import TransactionDraft
from core.i18n import tr, N_

HEADERS = (N_("Fecha"), N_("Cuenta"), N_("Concepto"), N_("Categoría"), N_("Importe"), N_("Saldo"))
SUPPORTED = (".csv", ".xlsx")


class ExportError(ValueError):
    """Formato no admitido o archivo que no se puede escribir."""


def _spanish(value: Decimal | None) -> str:
    return "" if value is None else f"{value:.2f}".replace(".", ",")


def write_movements(
    path: str | Path,
    movements: Sequence[TransactionDraft],
    account_names: Mapping[int | None, str],
    category_name,
) -> int:
    """Escribe los movimientos en el orden recibido. El formato lo decide la extension.
    Devuelve cuantos ha escrito. Los importes van como numeros en Excel y con coma decimal en CSV."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ExportError(tr("Solo se puede exportar a CSV o Excel (.csv, .xlsx)."))
    rows = [
        (t.booking_date, account_names.get(t.account_id, ""), t.description, category_name(t.category), t.amount, t.balance_after)
        for t in movements
    ]
    try:
        if suffix == ".csv":
            with open(path, "w", newline="", encoding="utf-8-sig") as f:      # utf-8-sig: Excel lo abre con acentos
                writer = csv.writer(f, delimiter=";")
                writer.writerow([tr(h) for h in HEADERS])
                for day, account, concept, category, amount, balance in rows:
                    writer.writerow([f"{day:%d/%m/%Y}", account, concept, category, _spanish(amount), _spanish(balance)])
        else:
            from openpyxl import Workbook                   # solo se carga al exportar a Excel (tarda unos 300 ms)
            from openpyxl.styles import Font

            book = Workbook()
            sheet = book.active
            sheet.title = tr("Movimientos")
            sheet.append([tr(h) for h in HEADERS])
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for day, account, concept, category, amount, balance in rows:
                sheet.append([day, account, concept, category, float(amount), None if balance is None else float(balance)])
            for r in range(2, len(rows) + 2):
                sheet.cell(r, 1).number_format = "DD/MM/YYYY"
                sheet.cell(r, 5).number_format = sheet.cell(r, 6).number_format = '#,##0.00 "€"'
            for column, width in zip("ABCDEF", (12, 16, 44, 24, 14, 14)):
                sheet.column_dimensions[column].width = width
            sheet.freeze_panes = "A2"
            book.save(path)
    except OSError as exc:
        raise ExportError(tr("No se pudo guardar el archivo: {exc}").format(exc=exc.strerror or exc)) from exc
    return len(rows)
