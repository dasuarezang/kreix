"""Funciones nuevas de la interfaz: categorias editables, presupuestos, pagos recurrentes, exportar y estado."""
from datetime import date
from decimal import Decimal as D
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QMessageBox

from core.enums import TxCategory
from services.movements_export import ExportError, write_movements
from ui.format import category_label
from ui.views.accounts_view import delta_text

@pytest.fixture(autouse=True)
def no_modal_dialogs(monkeypatch):
    monkeypatch.setattr(QMessageBox, "warning", lambda *a, **k: None)
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)


@pytest.fixture()
def loaded(window, tmp_path):
    """Cinco cargos mensuales de Netflix y un Mercadona, ya importados."""
    rows = ["Concepto;Fecha;Importe;Saldo"]
    balance = 500.0
    lines = []
    for month in range(5, 10):
        balance -= 12.99
        lines.append((f"05/{month:02d}/2026", "NETFLIX.COM", -12.99, balance))
    balance -= 20.5
    lines.append(("08/09/2026", "MERCADONA CL TEST", -20.5, balance))
    for day, concept, amount, bal in reversed(lines):
        rows.append(f"{concept};{day};{amount:.2f}EUR;{bal:.2f}EUR".replace(f"{amount:.2f}", f"{amount:.2f}".replace(".", ","))
                    .replace(f"{bal:.2f}", f"{bal:.2f}".replace(".", ",")))
    f = tmp_path / "extractImagin"
    f.write_text("\n".join(rows) + "\n", encoding="utf-8")
    window.import_view.import_paths([str(f)])
    return window


def movement(view, text):
    return next(t for t in view.table.visible_rows() if text in t.description)


def test_recurring_panel_lists_the_subscription_with_its_yearly_cost(loaded):
    panel = loaded.accounts_view.recurring_panel
    assert [r.name for r in panel.table.visible_rows()] == ["NETFLIX.COM"]
    assert "12,99" in panel.subtitle.text() and "155,88" in panel.subtitle.text()     # 12,99 x 12


def test_changing_a_category_by_hand_updates_the_table_and_survives_reload(loaded):
    v = loaded.accounts_view
    v.change_category(movement(v, "MERCADONA"), TxCategory.LEISURE)
    assert movement(v, "MERCADONA").category == TxCategory.LEISURE and movement(v, "MERCADONA").is_manual
    loaded.refresh()
    assert movement(v, "MERCADONA").category == TxCategory.LEISURE
    v.change_category(movement(v, "MERCADONA"), None)
    assert movement(v, "MERCADONA").category == TxCategory.GROCERIES and not movement(v, "MERCADONA").is_manual


def test_a_rule_created_from_the_view_recategorizes_and_shows_in_settings(loaded):
    v = loaded.accounts_view
    assert v.add_rule("netflix", TxCategory.LEISURE)
    assert {t.category for t in v.table.visible_rows() if "NETFLIX" in t.description} == {TxCategory.LEISURE}
    rules = loaded.settings_view.rules
    assert [(r.keyword, r.category) for r in rules._service.rules()] == [("NETFLIX", TxCategory.LEISURE)]
    assert rules._rows.count() == 1                                   # la lista de Ajustes se ha actualizado
    rules.remove(rules._service.rules()[0].id)
    assert {t.category for t in v.table.visible_rows() if "NETFLIX" in t.description} == {TxCategory.SUBSCRIPTIONS}


def test_budgets_show_progress_and_flag_the_ones_exceeded(loaded):
    v = loaded.accounts_view
    assert v.budget_panel.edit_button.text() == "Definir presupuestos" and v.budget_panel.bars.isHidden()
    v.save_budgets({TxCategory.GROCERIES: D(3), TxCategory.SUBSCRIPTIONS: D(100)})
    lines = v.budget_panel.bars._lines
    assert {(l.category, l.state) for l in lines} == {(TxCategory.GROCERIES, "over"), (TxCategory.SUBSCRIPTIONS, "ok")}
    assert v.budget_panel.edit_button.text() == "Editar" and "1 superado" in v.budget_panel.subtitle.text()
    v.period.select_month(date(2026, 9, 1))                         # en un mes: gasto de ese mes
    assert {l.category: l.spent for l in v.budget_panel.bars._lines}[TxCategory.GROCERIES] == D("20.5")


def test_month_cards_compare_with_the_previous_month(loaded):
    v = loaded.accounts_view
    v.period.select_month(date(2026, 9, 1))
    assert "frente a ago 26" in v.card_expenses._sub.text()          # 33,49 frente a 12,99
    v.period.select_month(date(2026, 5, 1))
    assert "frente a" not in v.card_expenses._sub.text()             # el primer mes no tiene anterior


def test_delta_text():
    aug = date(2026, 8, 1)
    assert delta_text(D(112), D(100), aug) == "+12 % frente a ago 26"
    assert delta_text(D(75), D(100), aug).startswith("−25 %")
    assert delta_text(D("100.2"), D(100), aug) == "Igual que en ago 26"
    assert delta_text(D(5), D(0), aug) == ""


def test_export_writes_exactly_the_visible_rows_to_csv_and_excel(loaded, tmp_path):
    v = loaded.accounts_view
    v.search.setText("netflix")                                        # solo se exporta lo filtrado
    csv_path, xlsx_path = tmp_path / "m.csv", tmp_path / "m.xlsx"
    assert v.export_movements(str(csv_path)) == 5
    text = csv_path.read_text(encoding="utf-8-sig").splitlines()
    assert text[0] == "Fecha;Cuenta;Concepto;Categoría;Importe;Saldo"
    assert text[1].startswith("05/09/2026;") and ";-12,99;" in text[1] and len(text) == 6
    assert v.export_movements(str(xlsx_path)) == 5
    sheet = load_workbook(xlsx_path).active
    assert sheet.cell(2, 5).value == pytest.approx(-12.99) and sheet.cell(2, 3).value == "NETFLIX.COM"
    assert sheet.cell(2, 4).value == category_label(TxCategory.SUBSCRIPTIONS)


def test_export_refuses_pdf_and_empty_selection(loaded, tmp_path):
    with pytest.raises(ExportError):
        write_movements(tmp_path / "x.pdf", [], {}, str)
    v = loaded.accounts_view
    v.search.setText("zzz-no-existe")
    assert v.export_movements(str(tmp_path / "vacio.csv")) is None
    assert not (tmp_path / "vacio.csv").exists()


def test_shortcuts_focus_the_search_and_reload_the_page(loaded):
    loaded.show()
    loaded.navigate(loaded.SUMMARY)
    loaded.focus_search()
    assert loaded.stack.currentIndex() == loaded.ACCOUNTS
    assert loaded.accounts_view.search.hasFocus() or loaded.accounts_view.search.isVisible()


def test_window_remembers_size_and_last_page(tmp_path, qapp):
    from main import build_window
    from tests.conftest import dispose_windows

    url = f"sqlite:///{tmp_path / 'k.db'}"
    first = build_window(url, "k")
    first.resize(1333, 777)
    first.navigate(first.ACCOUNTS)
    first.close()
    first.deleteLater()
    second = build_window(url, "k")
    try:
        state = second._svc.settings.load_state()
        assert (state["width"], state["height"], state["page"]) == ("1333", "777", str(second.ACCOUNTS))
        assert second.stack.currentIndex() == second.ACCOUNTS
        assert second._nav.button(second.ACCOUNTS).isChecked()
    finally:
        dispose_windows(second)


def test_summary_shows_upcoming_payments_and_all_the_budgets(loaded):
    d = loaded.dashboard
    assert [r.name for r in d.upcoming.table.visible_rows()] == ["NETFLIX.COM"]
    assert "Sin presupuestos" in d.budgets.subtitle.text() and d.budgets.lines == []
    loaded.accounts_view.save_budgets({TxCategory.GROCERIES: D(10), TxCategory.SUBSCRIPTIONS: D(500)})
    loaded.refresh()
    assert [(l.category, l.state) for l in d.budgets.lines] == [(TxCategory.GROCERIES, "over"), (TxCategory.SUBSCRIPTIONS, "ok")]
    assert "2 presupuestos" in d.budgets.subtitle.text() and "1 superado" in d.budgets.subtitle.text()
    loaded.accounts_view.save_budgets({TxCategory.GROCERIES: D(500)})
    loaded.refresh()
    assert [(l.category, l.state) for l in d.budgets.lines] == [(TxCategory.GROCERIES, "ok")]     # uno solo: sin flechas ni puntos
    assert d.budgets.dots.isHidden() and d.budgets._controls.isHidden()


def test_info_buttons_open_their_help_entry(window):
    from ui.widgets.help_link import InfoButton
    buttons = window.dashboard.findChildren(InfoButton) + window.accounts_view.findChildren(InfoButton) + window.investment_view.findChildren(InfoButton)
    assert {b.term for b in buttons} >= {"Patrimonio neto", "Tasa de ahorro", "Saldo", "Valor de la cartera", "Plusvalía y minusvalía"}
    next(b for b in buttons if b.term == "Patrimonio neto").click()
    assert window.stack.currentIndex() == window.HELP and "Patrimonio neto" in window.help_view.entry_terms()
    from core import i18n
    from ui.help_content import ENTRIES, localized_term
    assert all(localized_term(e.term) for e in ENTRIES) and all(b.term in {e.term for e in ENTRIES} for b in buttons)


def test_home_shows_the_net_worth_change_against_last_month(window):
    window.start_demo()
    data = window._svc.dashboard.load()
    assert data.net_worth_change == data.wealth[-1].net_worth - data.wealth[-2].net_worth
    chips = [c.text() for c in window.dashboard.card_net._chip_widgets]
    assert any("frente al mes pasado" in t for t in chips)


def test_the_user_can_mark_and_unmark_a_recurring_payment_from_a_movement(loaded):
    window = loaded
    view = window.accounts_view
    mercadona = next(t for t in view._data.movements if "MERCADONA" in t.description)
    offer = view.recurring_action(mercadona)
    assert offer is not None and "Marcar" in offer[0] and offer[1] is True                  # Kreix no lo detecta: se ofrece marcarlo
    assert "MERCADONA" not in [r.name for r in view._data.recurring]
    assert view.set_recurring(mercadona.description, True)
    assert any("MERCADONA" in r.name and r.manual for r in view._data.recurring)
    shown = [view.recurring_panel.table.cell_texts(i)[0] for i in range(view.recurring_panel.table.proxy.rowCount())]
    assert any("MERCADONA" in t for t in shown) and not any("marcado por ti" in t for t in shown)      # no se distingue el origen
    text, include = view.recurring_action(mercadona)
    assert "no es un pago recurrente" in text and include is False                                   # un pago marcado tambien se puede quitar
    window.dashboard.set_data(window._svc.dashboard.load())                                  # y tambien en Proximos pagos
    assert any("MERCADONA" in r.name for r in window._svc.dashboard.load().upcoming)
    assert view.set_recurring(mercadona.description, None)
    assert "MERCADONA" not in [r.name for r in view._data.recurring]


def test_the_user_can_say_a_detected_payment_is_not_recurring(loaded):
    view = loaded.accounts_view
    netflix = next(t for t in view._data.movements if "NETFLIX" in t.description)
    text, include = view.recurring_action(netflix)
    assert "no es un pago recurrente" in text and include is False
    view.set_recurring(netflix.description, False)
    assert not [r for r in view._data.recurring if "NETFLIX" in r.name]
    assert not [r for r in loaded._svc.dashboard.load().upcoming if "NETFLIX" in r.name]
    assert "Volver a detectar" in view.recurring_action(netflix)[0]


def test_net_worth_shows_liquidity_and_investment_with_their_share(window):
    window.start_demo()
    parts = window.dashboard.card_net.split._parts
    assert [p[0] for p in parts] == ["Liquidez", "Inversión"]
    assert all("€" in p[1] and "%" in p[1] for p in parts)                                    # importe y porcentaje del patrimonio
    shares = [float(p[1].split("·")[-1].replace("%", "").replace(",", ".").strip()) for p in parts]
    assert abs(sum(shares) - 100) < 0.2


def test_any_recurring_payment_can_be_removed_from_upcoming_payments(loaded):
    window = loaded
    upcoming = window.dashboard.upcoming
    netflix = next(r for r in window._svc.dashboard.load().upcoming if "NETFLIX" in r.name)          # detectado por Kreix
    upcoming.remove_requested.emit(netflix)
    assert not [r for r in window._svc.dashboard.load().upcoming if "NETFLIX" in r.name]
    assert not [r for r in window.accounts_view._data.recurring if "NETFLIX" in r.name]
